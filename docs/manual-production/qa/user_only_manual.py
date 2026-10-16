from docx import Document
from pathlib import Path
from docx.oxml.ns import qn
import re

root=Path(r'D:\EIS-DOCS\docs\manual-production')
path=root/'EIS-DOCS_Production_User_Manual_TH.docx'
doc=Document(path)
body=doc.element.body
def text(node): return ''.join(node.itertext()) if False else ''.join(node.xpath('.//w:t/text()'))
def heading(node):
 ppr=node.find(qn('w:pPr'))
 style=ppr.find(qn('w:pStyle')) if ppr is not None else None
 return style is not None and doc.styles[style.get(qn('w:val'))].name in ['Heading 1','Heading 2']
def remove_section(prefix):
 nodes=list(body)
 start=next(i for i,n in enumerate(nodes) if heading(n) and text(n).startswith(prefix))
 level=nodes[start].find(qn('w:pPr')).find(qn('w:pStyle')).get(qn('w:val'))
 end=next((i for i in range(start+1,len(nodes)) if heading(nodes[i]) and nodes[i].find(qn('w:pPr')).find(qn('w:pStyle')).get(qn('w:val'))==level),len(nodes)-1)
 if doc.styles[level].name=='Heading 1' and start>0 and nodes[start-1].xpath('.//w:br[@w:type="page"]'):start-=1
 if doc.styles[level].name=='Heading 1' and end>0 and nodes[end-1].xpath('.//w:br[@w:type="page"]'):end-=1
 for node in nodes[start:end]:body.remove(node)

for prefix in ['23 งาน Admin','24 งาน Admin','26 ภาคผนวกผู้ดูแล']:remove_section(prefix)
remove_section('เพิ่มชนิดเอกสาร เฉพาะ Admin')
# Remove administrative role and menu rows, preserving user support references.
for table in doc.tables:
 for row in list(table.rows):
  if row.cells[0].text.strip() in ['Admin','เมนู Admin']:table._tbl.remove(row._tr)

replacements={
 'ฉบับ 1.0  |  วันที่ 7 ตุลาคม 2569':'ฉบับ 1.1  |  วันที่ 7 ตุลาคม 2569',
 'สำหรับการใช้งานบน Production':'สำหรับผู้ใช้งานระบบ Production',
 'สำหรับการใช้งานบน Production':'สำหรับผู้ใช้งานระบบ Production',
 'ผู้จัดทำเอกสาร ผู้ตรวจสอบ ผู้ลงนาม ผู้อนุมัติ และ Admin':'ผู้จัดทำเอกสาร ผู้ตรวจสอบ ผู้ลงนาม และผู้อนุมัติ',
 'อ้างอิงโค้ด EIS-DOCS ที่ตรวจวันที่ 7 ตุลาคม 2569 และหน้าจอที่ระบุในคู่มือ':'ครอบคลุมการใช้งานเอกสารของผู้ใช้ตามหน้าที่และสิทธิ์ที่ได้รับ',
 'ผู้ดูแลต้องตรวจว่ารุ่นบน Production ตรงกับฟังก์ชันในคู่มือนี้ก่อนเผยแพร่ให้ผู้ใช้':'ใช้คู่มือนี้ตามเมนูที่บัญชีได้รับสิทธิ์ หากหน้าจอต่างจากคู่มือให้ติดต่อผู้ดูแลระบบ',
 'คู่มือแบ่งขั้นตอนให้ตรงกับเมนูและชื่อปุ่มในระบบ':'คู่มือแบ่งขั้นตอนให้ตรงกับเมนูและชื่อปุ่มที่ผู้ใช้งานพบในระบบ',
 'เลือกเมนูผู้ใช้บริเวณชื่อบัญชี':'เลือกเมนูผู้ใช้บริเวณชื่อบัญชี',
 'การเป็นผู้สร้างไฟล์หรืออยู่หน่วยงานเดียวกันไม่ได้ทำให้มีสิทธิ์แก้ไขทุกเอกสาร':'การอยู่หน่วยงานเดียวกันไม่ได้ทำให้มีสิทธิ์แก้ไขทุกเอกสาร',
 'ชื่อเอกสารและชนิด':'ชื่อเอกสารและชนิด',
 'Template ผังในหน้าวางตำแหน่ง':'Template ผังในหน้าวางตำแหน่ง',
}
for p in doc.paragraphs:
 value=p.text
 for a,b in replacements.items():value=value.replace(a,b)
 if value!=p.text:p.text=value
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    value=p.text
    for a,b in replacements.items():value=value.replace(a,b)
    if value!=p.text:p.text=value

# Update instructions to match current date and bookmark behavior.
for p in doc.paragraphs:
 if p.text.startswith('การแก้ไขใช้ไฟล์รุ่นใหม่'): pass
 if p.text.startswith('โดยทั่วไป WF1 ของเอกสาร'):
  p.text='WF1 ของเอกสารเป็นผู้แก้ไขตามสิทธิ์ที่ได้รับ ระบบกำหนดให้ WF1 ยังมีสถานะ PENDING และยังไม่ลงนาม หากลงนามแล้วและต้องแก้ไข ให้ใช้กระบวนการส่งกลับหรือประสานผู้ดูแลตามวิธีปฏิบัติของหน่วยงาน'
 if p.text.startswith('เมื่อเปลี่ยนชื่อปุ่ม สิทธิ์ หรือ Workflow'):
  p.text='เมื่อหน้าจอหรือขั้นตอนเปลี่ยน ให้ใช้คู่มือฉบับล่าสุดจากหน่วยงาน หากขั้นตอนในคู่มือไม่ตรงกับหน้าจอ ให้แจ้งผู้ดูแลพร้อมชื่อเมนูและข้อความที่พบ'
 if p.text.startswith('5. เลือกวิธีแสดงวันที่ลงนาม'):
  p.text='5. เลือกวิธีแสดงวันที่ลงนาม ค่าเริ่มต้นเอกสารใหม่คือไม่แสดงวันที่ลงนาม หากเลือกโหมดนี้ หน้าวางตำแหน่ง ตัวอย่าง และ PDF จะไม่แสดงวันที่ แม้ดำเนินการลงนามแล้ว'
 if p.text.startswith('4. อ่าน PDF ทุกหน้าที่เกี่ยวข้อง'):
  p.text='4. อ่าน PDF ทุกหน้าที่เกี่ยวข้อง คลิกชื่อ Bookmark ทางขวาเพื่อเปิด PDF ไปยังเลขหน้าที่ระบุ รายการที่เลือกจะแสดงไฮไลต์'
 if p.text.startswith('1. ดูรายการบุ๊คมาร์ค') or p.text.startswith('1. ตรวจรายการบุ๊คมาร์ค'):
  p.text='1. กดโหลด Bookmark จากข้อมูลหลัก เพื่อดึงชื่อและเลขหน้าล่าสุดที่เลือกในข้อมูลหลัก หรือใช้รายการบุ๊คมาร์คจาก PDF ตามงานที่จัดทำ'

for table in doc.tables:
 if table.cell(0,0).text=='วิธีวันที่':
  row=table.add_row();row.cells[0].text='ไม่แสดงวันที่ลงนาม';row.cells[1].text='ค่าเริ่มต้นเอกสารใหม่ ไม่แสดงวันที่ในผัง ตัวอย่าง และ PDF หลังลงนาม'
 if table.cell(0,0).text=='อาการ':
  for row in table.rows:
   if row.cells[0].text=='Production ต่างจาก Local':
    row.cells[0].text='หน้าจอไม่ตรงกับคู่มือ';row.cells[1].text='รีเฟรชหน้าเว็บ หากยังต่างให้แจ้งผู้ดูแลพร้อมชื่อเมนูและภาพหน้าจอ'

number=0
for p in doc.paragraphs:
 if p.style.name=='Heading 1' and re.match(r'^\d+ ',p.text):
  number+=1;p.text=re.sub(r'^\d+ ',f'{number} ',p.text)
for s in doc.sections:
 for p in s.footer.paragraphs:
  for run in p.runs:run.text=run.text.replace('ฉบับ 1.0','ฉบับ 1.1')
doc.core_properties.subject='คู่มือสำหรับผู้ใช้งานระบบ'
doc.core_properties.title='คู่มือผู้ใช้งานระบบ EIS-DOCS Production ภาษาไทย'
doc.save(path)
print(f'Saved user-only manual with {number} chapters: {path}')
