from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pathlib import Path
path=Path(r'D:\EIS-DOCS\docs\manual-production\EIS-DOCS_Production_User_Manual_TH.docx')
d=Document(path)
for p in d.paragraphs:
 if p.text.startswith('แจ้งเตือนช่วยติดตามงาน'):
  p.text='แจ้งเตือนช่วยติดตามงาน แต่ให้ยึดสถานะและลำดับจริงในรายการเอกสารเป็นหลัก หากไม่พบแจ้งเตือน ให้เปิดคิวและโหลดข้อมูลใหม่เพื่อตรวจงานของตนเอง'
 if p.text.startswith('4. หากเห็นแถบกำลังสวมสิทธิ์') or p.text.startswith('ขณะสวมสิทธิ์ไม่สามารถตั้ง'):
  p._element.getparent().remove(p._element)
 if p.text.startswith('5. เมื่อเสร็จงานเลือกออกจากระบบ'):
  p.text=p.text.replace('5.','4.',1)
for s in d.sections:
 p=s.footer.paragraphs[0];p.clear();p.alignment=2
 p.add_run('คู่มือฉบับ 1.1   •   หน้า ')
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
 for r in p.runs:r.font.size=Pt(8)
d.save(path)
print('Finalized user scope and page fields')
